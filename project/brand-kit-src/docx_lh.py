from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
P='/tmp/claude-0/kitwork/kit/03-print-and-documents'
LOGO='/tmp/claude-0/kitwork/kit/01-logos/lockup-horizontal-for-light-backgrounds@2x.png'
d=Document(); s=d.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11); s.left_margin=s.right_margin=Inches(0.9); s.top_margin=Inches(1.7); s.bottom_margin=Inches(1); s.header_distance=Inches(0.45); s.footer_distance=Inches(0.4)
st=d.styles['Normal']; st.font.name='IBM Plex Sans'; st.font.size=Pt(11); st.font.color.rgb=RGBColor(0x16,0x1a,0x22)
st.element.rPr.rFonts.set(qn('w:eastAsia'),'IBM Plex Sans')
st.paragraph_format.space_after=Pt(8); st.paragraph_format.line_spacing=1.25
# header: logo left, contact right (2-col borderless table)
h=s.header; t=h.add_table(rows=1,cols=2,width=Inches(6.7)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
c0,c1=t.rows[0].cells; c0.width=Inches(3.7); c1.width=Inches(3.0)
c0.paragraphs[0].add_run().add_picture(LOGO,height=Inches(0.75))
p=c1.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for i,line in enumerate(['www.redass.net','news.redass.net','discord.gg/gH577r59C8']):
    r=p.add_run(line+('\n' if i<2 else '')); r.font.size=Pt(9); r.font.color.rgb=RGBColor(0x56,0x60,0x70)
h.paragraphs[0].text=''
# red rule under header
hp=h.add_paragraph(); pPr=hp._p.get_or_add_pPr(); b=OxmlElement('w:pBdr'); bt=OxmlElement('w:bottom')
for k,v in (('w:val','single'),('w:sz','12'),('w:space','1'),('w:color','C4211B')): bt.set(qn(k),v)
b.append(bt); pPr.append(b)
# footer
f=s.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
for txt,col in (('DRAW. ','566070'),('BURN. ','C4211B'),('REPEAT.','566070'),('     ·     ADULTS 18+','566070')):
    r=f.add_run(txt); r.font.name='Oswald'; r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(col)
# body placeholder
d.add_paragraph('Date')
for t_ in ['','Dear [Name],','','Start typing here. The logo at the top and the line at the bottom repeat on every page.','','Cheers,','','[Your name]','Red Ass']:
    d.add_paragraph(t_)
d.save(f'{P}/letterhead-editable.docx'); print('ok')
